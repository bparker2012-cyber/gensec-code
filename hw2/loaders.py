"""Custom LangChain loader for the file types accepted by SourceLens."""

from __future__ import annotations

import csv
import json
from collections.abc import Iterator
from pathlib import Path

from docx import Document as WordDocument
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from pypdf import PdfReader


class NotebookFileLoader(BaseLoader):
    """Load common notebook files into normalized LangChain documents.

    Unlike a collection of unrelated pre-built loaders, this custom loader
    applies one validation policy and consistent source metadata across text,
    Markdown, CSV, JSON, PDF, and DOCX files.
    """

    SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}

    def __init__(self, file_path: str | Path, *, max_bytes: int = 10_000_000) -> None:
        """Store and validate the path that will be loaded lazily."""
        self.file_path = Path(file_path)
        self.max_bytes = max_bytes

        if max_bytes <= 0:
            raise ValueError("max_bytes must be greater than zero")
        if not self.file_path.is_file():
            raise FileNotFoundError(f"File does not exist: {self.file_path}")
        if self.file_path.suffix.lower() not in self.SUPPORTED_SUFFIXES:
            supported = ", ".join(sorted(self.SUPPORTED_SUFFIXES))
            raise ValueError(f"Unsupported file type. Choose one of: {supported}")
        self._validate_size()

    def lazy_load(self) -> Iterator[Document]:
        """Yield normalized documents with stable source metadata."""
        # BaseLoader is lazy, so the file may have changed since construction.
        self._validate_size()
        suffix = self.file_path.suffix.lower()
        if suffix in {".txt", ".md"}:
            yield self._document(self.file_path.read_text(encoding="utf-8"))
        elif suffix == ".csv":
            yield from self._load_csv()
        elif suffix == ".json":
            yield self._document(self._load_json())
        elif suffix == ".pdf":
            yield from self._load_pdf()
        elif suffix == ".docx":
            yield self._document(self._load_docx())

    def _validate_size(self) -> None:
        """Enforce the upload limit at construction and again at read time."""
        if self.file_path.stat().st_size > self.max_bytes:
            raise ValueError(f"File exceeds the {self.max_bytes:,}-byte limit")

    def _document(self, content: str, **metadata: object) -> Document:
        """Create one non-empty document using the shared metadata schema."""
        normalized = content.strip()
        if not normalized:
            raise ValueError(f"No readable text found in {self.file_path.name}")
        return Document(
            page_content=normalized,
            metadata={"source": self.file_path.name, "file_type": self.file_path.suffix.lower(), **metadata},
        )

    def _load_csv(self) -> Iterator[Document]:
        """Represent each CSV row as a labeled, independently retrievable document."""
        with self.file_path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                raise ValueError(f"CSV has no header row: {self.file_path.name}")
            headers = [header.strip() if header is not None else "" for header in reader.fieldnames]
            if any(not header for header in headers) or len(set(headers)) != len(headers):
                raise ValueError(f"CSV headers must be non-empty and unique: {self.file_path.name}")
            reader.fieldnames = headers
            yielded = False
            for row_number, row in enumerate(reader, start=2):
                if None in row:
                    raise ValueError(
                        f"CSV row {row_number} has more values than the header: "
                        f"{self.file_path.name}"
                    )
                values = {key: (value or "").strip() for key, value in row.items()}
                if not any(values.values()):
                    continue
                text = "\n".join(f"{key}: {value}" for key, value in values.items())
                yielded = True
                yield self._document(text, row=row_number)
            if not yielded:
                raise ValueError(f"CSV has no data rows: {self.file_path.name}")

    def _load_json(self) -> str:
        """Parse JSON before converting it into readable indented text."""
        with self.file_path.open(encoding="utf-8") as handle:
            data = json.load(handle)
        return json.dumps(data, indent=2, ensure_ascii=False)

    def _load_pdf(self) -> Iterator[Document]:
        """Yield one document per PDF page that contains extractable text."""
        reader = PdfReader(str(self.file_path))
        yielded = False
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                yielded = True
                yield self._document(text, page=page_number)
        if not yielded:
            raise ValueError(
                f"No extractable text found in {self.file_path.name}; scanned PDFs require OCR"
            )

    def _load_docx(self) -> str:
        """Extract paragraphs and table rows from a Word document."""
        document = WordDocument(str(self.file_path))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    blocks.append(" | ".join(values))
        return "\n".join(blocks)
